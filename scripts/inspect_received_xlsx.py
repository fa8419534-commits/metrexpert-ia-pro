from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import json
import openpyxl

path = Path('/home/ubuntu/upload/metrexpert-1787850043858.xlsx')
report = {'file': str(path), 'size_bytes': path.stat().st_size, 'sheets': [], 'formula_xml': [], 'warnings': []}

wb = openpyxl.load_workbook(path, data_only=False)
for ws in wb.worksheets:
    formulas = []
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith('='):
                formulas.append({'cell': cell.coordinate, 'formula': cell.value, 'number_format': cell.number_format})
    report['sheets'].append({
        'title': ws.title,
        'max_row': ws.max_row,
        'max_column': ws.max_column,
        'freeze_panes': str(ws.freeze_panes) if ws.freeze_panes else None,
        'print_area': str(ws.print_area) if ws.print_area else None,
        'page_orientation': ws.page_setup.orientation,
        'fit_to_width': ws.page_setup.fitToWidth,
        'fit_to_height': ws.page_setup.fitToHeight,
        'formula_count': len(formulas),
        'formulas': formulas,
        'merged_ranges': [str(x) for x in ws.merged_cells.ranges],
    })

ns = {'x': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
with ZipFile(path) as z:
    names = z.namelist()
    report['zip_entries'] = len(names)
    for name in names:
        if name.startswith('xl/worksheets/sheet') and name.endswith('.xml'):
            root = ET.fromstring(z.read(name))
            for f in root.findall('.//x:f', ns):
                text = f.text or ''
                report['formula_xml'].append({'xml_file': name, 'formula_text': text, 'has_internal_equals': text.startswith('=')})
    report['formula_internal_equals'] = [x for x in report['formula_xml'] if x['has_internal_equals']]
    report['media_entries'] = [x for x in names if x.startswith('xl/media/')]

if report['formula_internal_equals']:
    report['warnings'].append('Des formules XML contiennent encore un signe égal interne.')
if not report['media_entries']:
    report['warnings'].append('Aucun média incorporé ; attendu si aucune signature/tampon n’a été fourni.')

out = Path('/home/ubuntu/metrexpert-ia-pro/docs/inspection-metrexpert-1787850043858.json')
out.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({
    'output': str(out),
    'sheets': [(x['title'], x['formula_count'], x['max_row'], x['max_column']) for x in report['sheets']],
    'formula_xml_count': len(report['formula_xml']),
    'internal_equals': len(report['formula_internal_equals']),
    'media_entries': report['media_entries'],
    'warnings': report['warnings'],
}, ensure_ascii=False))
