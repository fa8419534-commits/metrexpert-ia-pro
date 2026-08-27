from pathlib import Path
import json
import openpyxl

path = Path('/home/ubuntu/upload/metrexpert-1787850043858.xlsx')
wb_formula = openpyxl.load_workbook(path, data_only=False)
wb_values = openpyxl.load_workbook(path, data_only=True)
report = {}
for title in wb_formula.sheetnames:
    wf, wv = wb_formula[title], wb_values[title]
    report[title] = {
        'formula_cells': {c.coordinate: c.value for row in wf.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('=')},
        'cached_values': {coord: wv[coord].value for coord in ['B15', 'F13', 'F2', 'D2'] if coord in wv},
        'tab_color': wf.sheet_properties.tabColor.rgb if wf.sheet_properties.tabColor and wf.sheet_properties.tabColor.type == 'rgb' else None,
        'print_title_rows': str(wf.print_title_rows) if wf.print_title_rows else None,
        'sheet_view_show_grid_lines': wf.sheet_view.showGridLines,
    }
report['workbook_calc'] = {
    'fullCalcOnLoad': wb_formula.calculation.fullCalcOnLoad,
    'forceFullCalc': wb_formula.calculation.forceFullCalc,
    'calcMode': wb_formula.calculation.calcMode,
}
Path('/home/ubuntu/metrexpert-ia-pro/docs/inspection-values-metrexpert-1787850043858.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(report, ensure_ascii=False))
